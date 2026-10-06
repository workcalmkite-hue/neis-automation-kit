# -*- coding: utf-8 -*-
"""근태 색인 목록(한글 .hwp) 채우기 — 구글 출결 캘린더에서 그 달 것을 골라 날짜순으로 끼워 넣는다.

  python geuntae.py 2026-10                 무엇을 넣을지 보기만 한다 (파일은 안 고친다)
  python geuntae.py 2026-10 --write         실제로 쓴다 (먼저 이 폴더 출력/백업/ 에 사본을 둔다)
  python geuntae.py 2026-10 --pages         그 달 줄이 있는 쪽 번호만 알려 준다 (인쇄는 print_month.py)
  python geuntae.py --hwp "파일경로"        근태 색인 목록 파일을 지정한다 (설정에 기억)

넣는 것 / 빼는 것 (방배중 근태신고서철 기준)
  넣는다  질병결석 · 출석인정결석 · 출석인정조퇴 · 출석인정지각
  뺀다    질병조퇴 · 질병지각 · 미인정(결석·지각·조퇴) 전부

칸 채우는 규칙
  질병결석                         질병결석 칸 O
  출석인정결석 + 사유(생리·조부상…)   출석인정 칸 «O (사유)»
  출석인정결석 + 사유 없음/체험학습   교외체험학습 — 기간(일) · 누적(일)만. 출석인정 칸은 비운다
  출석인정조퇴 / 출석인정지각        출석인정 칸 O + 비고 «출석인정조퇴 (교시)»

기간(일)은 나이스에서 받은 «월별 출결현황» 엑셀(출력/YYYY-MM_월별출결현황.xlsx)에 실제로 찍힌 날 수로 센다.
엑셀이 없으면 평일 − 공휴일(구글 «대한민국의 휴일» 캘린더)로 세고 «추정»이라고 표시한다 — 재량휴업일은 못 뺀다.

★ 한글 파일은 한글 프로그램(pyhwpx)으로만 고친다. 한글 MCP의 바꾸기·채우기는 .hwp 를 깨뜨린다.
★ 화면에는 학생 이름을 «박○예» 처럼 가려서 찍는다. 한글 파일에는 이름 전체가 들어간다.
마지막 줄에 항상 RESULT_JSON {...} 을 찍는다.
"""
import calendar
import json
import os
import re
import shutil
import subprocess
import sys
from datetime import date, datetime, timedelta
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
HERE = Path(__file__).resolve().parent
OUT_DIR = HERE / "출력"

INCLUDE = ("질병결석", "출석인정결석", "출석인정조퇴", "출석인정지각")
HOLIDAY_CAL = "ko.south_korea#holiday@group.v.calendar.google.com"
ROWS_PER_TABLE = 25

RESULT = {"ok": False, "month": None, "hwp": None, "new": 0, "skipped": [], "warnings": [], "pages": []}


def log(*a):
    print(*a, flush=True)


def mask(name: str) -> str:
    name = (name or "").strip()
    if len(name) <= 1:
        return name
    if len(name) == 2:
        return name[0] + "○"
    return name[0] + "○" * (len(name) - 2) + name[-1]


def school_key(d: date):
    """학년도 순서(3월 → 이듬해 2월)로 정렬하는 열쇠."""
    return (d.year if d.month >= 3 else d.year - 1, d.month if d.month >= 3 else d.month + 12, d.day)


def md(d: date) -> str:
    return f"{d.month}/{d.day}"


# ──────────────────────────────────────────────────────────────
# 근태 색인 목록 파일 찾기
# ──────────────────────────────────────────────────────────────
def candidate_dirs():
    home = Path.home()
    dirs = [home / "Desktop", home / "바탕 화면", home / "Documents", home / "문서", home / "Downloads"]
    for od in home.glob("OneDrive*"):
        dirs += [od / "바탕 화면", od / "Desktop", od / "문서", od / "Documents"]
    return [d for d in dirs if d.is_dir()]


def find_hwp():
    hits = []
    for d in candidate_dirs():
        try:
            for p in d.rglob("*.hwp"):
                if "근태" in p.name and "색인" in p.name and "출력" not in p.parts[-2:-1]:
                    hits.append(p)
        except Exception:
            pass
        if len(hits) > 30:
            break
    seen, out = set(), []
    for p in hits:
        k = str(p.resolve()).lower()
        if k not in seen:
            seen.add(k)
            out.append(p)
    return out


def resolve_hwp(cfg, args):
    import teacher_config
    if "--hwp" in args:
        p = Path(args[args.index("--hwp") + 1].strip('"'))
        if not p.exists():
            log(f"❌ 파일이 없습니다: {p}")
            return None
        cfg["geuntae_hwp"] = str(p)
        teacher_config.save_config({k: v for k, v in cfg.items()})
        log(f"📄 근태 색인 목록 파일을 기억했습니다: {p.name}")
        return p
    saved = cfg.get("geuntae_hwp")
    if saved and Path(saved).exists():
        return Path(saved)
    hits = find_hwp()
    if len(hits) == 1:
        cfg["geuntae_hwp"] = str(hits[0])
        teacher_config.save_config(cfg)
        log(f"📄 근태 색인 목록 파일을 찾아 기억했습니다: {hits[0]}")
        return hits[0]
    if not hits:
        log("❌ 이름에 «근태»와 «색인»이 들어간 .hwp 파일을 바탕화면·문서·다운로드에서 못 찾았습니다.")
        log('   파일 위치를 알려 주세요:  python geuntae.py --hwp "파일경로"')
    else:
        log("❓ 근태 색인 목록 파일이 여러 개입니다. 쓸 파일을 골라 주세요:")
        for p in hits[:15]:
            log("   -", p)
        log('   python geuntae.py --hwp "파일경로"')
    RESULT["candidates"] = [str(p) for p in hits[:15]]
    return None


# ──────────────────────────────────────────────────────────────
# 구글 캘린더
# ──────────────────────────────────────────────────────────────
TITLE = re.compile(r"^\s*(\d+)\s*번?\s+([^\s(]+)(?:\s+([^\s(]+))?\s*(?:[(（]\s*(.*?)\s*[)）])?\s*$")


def ev_dates(ev):
    s, e = ev["start"], ev["end"]
    if "date" in s:
        a = date.fromisoformat(s["date"])
        b = date.fromisoformat(e["date"]) - timedelta(days=1)      # 종일 일정의 끝은 «다음 날»로 들어 있다
    else:
        a = datetime.fromisoformat(s["dateTime"]).date()
        b = datetime.fromisoformat(e["dateTime"]).date()
    return a, max(a, b)


def read_calendar(svc, cals, y, m):
    first = date(y, m, 1)
    last = date(y, m, calendar.monthrange(y, m)[1])
    tmin = (first - timedelta(days=1)).isoformat() + "T00:00:00Z"
    tmax = (last + timedelta(days=2)).isoformat() + "T00:00:00Z"
    out, bad = [], []
    for cat in INCLUDE:
        cid = cals.get(cat)
        if not cid:
            continue
        tok = None
        while True:
            r = svc.events().list(calendarId=cid, timeMin=tmin, timeMax=tmax, singleEvents=True,
                                  orderBy="startTime", pageToken=tok, maxResults=250).execute()
            for ev in r.get("items", []):
                a, b = ev_dates(ev)
                if not (first <= a <= last):          # 그 달에 «시작한» 것만. 전달에서 넘어온 건 전달 줄에 있다
                    continue
                t = (ev.get("summary") or "").strip()
                mt = TITLE.match(t)
                if not mt:
                    bad.append(f"{md(a)} {cat} «{t[:20]}»")
                    continue
                num, name, _cat, memo = mt.group(1), mt.group(2), mt.group(3), (mt.group(4) or "").strip()
                out.append({"cat": cat, "start": a, "end": b, "num": num, "name": name, "memo": memo})
            tok = r.get("nextPageToken")
            if not tok:
                break
    return out, bad


def public_holidays(svc, a: date, b: date) -> set:
    try:
        r = svc.events().list(calendarId=HOLIDAY_CAL, timeMin=a.isoformat() + "T00:00:00Z",
                              timeMax=(b + timedelta(days=1)).isoformat() + "T00:00:00Z",
                              singleEvents=True, maxResults=250).execute()
    except Exception:
        return set()
    days = set()
    for ev in r.get("items", []):
        s, e = ev_dates(ev)
        while s <= e:
            days.add(s)
            s += timedelta(days=1)
    return days


# ──────────────────────────────────────────────────────────────
# 나이스 월별 출결현황 엑셀 (실제로 찍힌 날)
# ──────────────────────────────────────────────────────────────
def neis_days(y, m):
    """{(번호, 출결구분): set(날짜)} — month_close.py --download 로 받은 엑셀. 없으면 None."""
    f = OUT_DIR / f"{y}-{m:02d}_월별출결현황.xlsx"
    if not f.exists():
        return None
    import warnings
    import openpyxl
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        ws = openpyxl.load_workbook(f).active
    # 같은 학생이 이어지면 번호·성명 칸이 병합돼 빈칸으로 읽힌다 → 병합 범위의 값으로 메운다
    fill = {}
    for rng in ws.merged_cells.ranges:
        vals = [ws.cell(r, c).value for r in range(rng.min_row, rng.max_row + 1)
                for c in range(rng.min_col, rng.max_col + 1)]
        v = next((x for x in vals if x not in (None, "")), None)
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                fill[(r, c)] = v
    head = [str(c.value or "").strip() for c in ws[1]]
    try:
        ci = {k: head.index(k) + 1 for k in ("일자", "번호", "출결구분")}
    except ValueError:
        return None
    out = {}
    for r in range(2, ws.max_row + 1):
        def val(k):
            v = ws.cell(r, ci[k]).value
            return fill.get((r, ci[k]), v) if v in (None, "") else v
        d, n, g = val("일자"), val("번호"), val("출결구분")
        mt = re.match(r"(\d{4})\.(\d{2})\.(\d{2})", str(d or ""))
        if not (mt and n and g):
            continue
        out.setdefault((str(n).strip(), str(g).strip()), set()).add(date(*map(int, mt.groups())))
    return out


# ──────────────────────────────────────────────────────────────
# 한글 표 읽기 · 쓰기 (pyhwpx = 한글 프로그램)
# ──────────────────────────────────────────────────────────────
def norm(s):
    return re.sub(r"\s+", "", str(s or ""))


def colmap(columns, row0):
    """머리글 글자로 칸 위치를 찾는다 (표마다 칸 수가 다르다 — 3·4쪽엔 «미인정결석» 칸이 하나 더 있다)."""
    cm = {}
    for i, (c, r) in enumerate(zip(columns, row0)):
        c, r = norm(c), norm(r)
        if "관련" in c:
            cm["no"] = i
        elif "근태" in c or "월/일" in c:
            cm["date"] = i
        elif c == "번호":
            cm["num"] = i
        elif "성명" in c:
            cm["name"] = i
        elif "질병" in c:
            cm["sick"] = i
        elif "미인정" in c:
            cm["unexcused"] = i
        elif "기타" in c:
            cm["etc"] = i
        elif "출석인정" in c:
            cm["excused"] = i
        elif "체험" in c and "기간" in r:
            cm["days"] = i
        elif "체험" in c and "누적" in r:
            cm["cum"] = i
        elif "비고" in c:
            cm["note"] = i
    need = {"date", "num", "name", "sick", "excused", "days", "cum", "note"}
    return cm if need <= set(cm) else None


def letter(i):
    return chr(ord("A") + i)


def _hwp_procs():
    """[(pid, 창제목)] — 창제목이 N/A 면 화면에 안 보이는 «남은» 한글이다 (지난번 자동화가 못 닫은 것)."""
    import csv
    import io
    r = subprocess.run(["tasklist", "/FI", "IMAGENAME eq Hwp.exe", "/V", "/FO", "CSV", "/NH"],
                       capture_output=True, text=True, encoding="mbcs", errors="ignore")
    out = []
    for row in csv.reader(io.StringIO(r.stdout or "")):
        if len(row) >= 9 and row[0].lower() == "hwp.exe":
            out.append((row[1], row[-1]))
    return out


def hwp_running():
    """선생님이 «화면에 띄워 둔» 한글이 있는가."""
    return any(t.strip() not in ("N/A", "", "OleMainThreadWndName") for _, t in _hwp_procs())


def kill_orphan_hwp():
    """화면에 안 보이는 한글만 끈다. 끈 게 있으면 2초 기다린다 (바로 새로 띄우면 죽어 가는 쪽에 붙는다)."""
    import time
    killed = 0
    for pid, t in _hwp_procs():
        if t.strip() in ("N/A", "", "OleMainThreadWndName"):
            subprocess.run(["taskkill", "/F", "/PID", pid], capture_output=True)
            killed += 1
    if killed:
        time.sleep(2)


def open_hwp(path: Path):
    import pyhwpx
    hwp = pyhwpx.Hwp(new=True, visible=False)
    hwp.open(str(path), "HWP", "")
    return hwp


def close_hwp(hwp):
    try:
        hwp.Clear(1)          # 저장 안 묻고 닫는다 (저장은 이미 save_as 로 했다)
    except Exception:
        pass
    try:
        hwp.Quit()
    except Exception:
        pass


def read_rows(hwp):
    """[{n, table, row, cm, cells{key:text}}] — 표 순서대로 모든 줄 (빈 줄 포함)."""
    rows, t = [], 0
    while True:
        try:
            df = hwp.table_to_df(t)
        except Exception:
            break
        cols = list(df.columns)
        row0 = [str(x) for x in df.iloc[0].tolist()]
        cm = colmap(cols, row0)
        if cm is None:
            t += 1
            continue
        for i in range(1, len(df)):
            vals = [str(x).replace("\r", " ").replace("\n", " ").strip() for x in df.iloc[i].tolist()]
            cells = {k: vals[ci] for k, ci in cm.items() if ci < len(vals)}
            rows.append({"table": t, "row": i + 2, "cm": cm, "cells": cells})
        t += 1
    for k, r in enumerate(rows, 1):
        r["n"] = k
    return rows


DATE1 = re.compile(r"(\d{1,2})\s*/\s*(\d{1,2})")


def parse_cell_date(s, school_year):
    ds = DATE1.findall(s or "")
    if not ds:
        return None, None
    def mk(mm, dd):
        mm, dd = int(mm), int(dd)
        return date(school_year if mm >= 3 else school_year + 1, mm, dd)
    a = mk(*ds[0])
    b = mk(*ds[-1]) if len(ds) > 1 else a
    return a, b


def kind_of(cells):
    """기존 줄이 무엇으로 표시됐는지."""
    if cells.get("days"):
        return "체험"
    if cells.get("sick"):
        return "질병결석"
    if cells.get("excused"):
        return "출석인정"
    if cells.get("etc") or cells.get("unexcused"):
        return "기타"
    return "?"


def entry_cells(e):
    """새 줄 하나의 칸 값 (cm 키 기준)."""
    c = {"date": e["date_str"], "num": e["num"], "name": e["name"]}
    if e["kind"] == "질병결석":
        c["sick"] = "O"
    elif e["kind"] == "출석인정":
        memo = re.sub(r"\s*결석$", "", e["memo"]) or e["memo"]      # «생리결석» → «생리» (서식 머리글이 «경조사/생리/공결»)
        c["excused"] = f"O ({memo})" if memo else "O"
    elif e["kind"] == "인정조퇴지각":
        c["excused"] = "O"
        c["note"] = f"{e['cat']} ({e['memo']})" if e["memo"] else e["cat"]
    elif e["kind"] == "체험":
        c["days"] = str(e["days"])
        c["cum"] = str(e["cum"])
    return c


def write_row(hwp, slot, cells):
    cm = slot["cm"]
    keys = [k for k in cm if k != "no"]
    for k in keys:                                      # 1) 줄 비우기 (관련번호 칸은 건드리지 않는다)
        hwp.get_into_nth_table(slot["table"])
        if hwp.goto_addr(f"{letter(cm[k])}{slot['row']}"):
            hwp.MoveSelListEnd()
            hwp.Delete()
    bad = []
    for k, v in cells.items():                          # 2) 값 쓰기
        if k not in cm or v in (None, ""):
            continue
        hwp.get_into_nth_table(slot["table"])           # goto_addr 는 셀마다 표에 다시 들어간 직후에만 믿을 수 있다
        if hwp.goto_addr(f"{letter(cm[k])}{slot['row']}"):
            hwp.insert_text(str(v))
            if k == "note" and len(str(v)) > 6:
                # 비고 칸은 좁다 — 보통 글씨로 쓰면 세 줄로 접히며 줄이 높아지고, 표가 다음 쪽으로 밀린다 (2026-10-06 실측)
                try:
                    hwp.MoveSelListBegin()
                    hwp.set_font(Height=7)
                    hwp.Cancel()
                except Exception:
                    pass
        else:
            bad.append(k)
    return bad


def dump_rows(path: Path):
    """다른 프로세스에서 한글로 열어 모든 줄을 읽어 온다."""
    r = subprocess.run([sys.executable, "-X", "utf8", str(Path(__file__).resolve()), "--_dump", str(path)],
                       capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=180)
    for line in (r.stdout or "").splitlines():
        if line.startswith("DUMP_JSON "):
            return json.loads(line[10:])
    return None


def _dump(path):
    kill_orphan_hwp()
    hwp = open_hwp(Path(path))
    try:
        rows = [{"n": r["n"], "cells": r["cells"]} for r in read_rows(hwp)]
    finally:
        close_hwp(hwp)
    print("DUMP_JSON " + json.dumps(rows, ensure_ascii=False))


# ──────────────────────────────────────────────────────────────
def main():
    args = sys.argv[1:]
    import teacher_config
    import main as M
    try:
        M.load_teacher_settings()
    except teacher_config.ConfigMissingError as e:
        log(f"❌ {e}")
        RESULT["error"] = "config"
        return 1
    cfg = teacher_config.load_config()
    path = resolve_hwp(cfg, args)
    ym = next((a for a in args if re.fullmatch(r"\d{4}-\d{1,2}", a)), None)
    if path is None or ym is None:
        if path is not None and ym is None:
            RESULT["ok"] = True
            RESULT["hwp"] = path.name
        return 0 if RESULT["ok"] else 1
    y, m = map(int, ym.split("-"))
    RESULT["month"] = f"{y}-{m:02d}"
    RESULT["hwp"] = path.name
    write = "--write" in args
    school_year = y if m >= 3 else y - 1

    kill_orphan_hwp()
    if hwp_running():
        log("⛔ 한글 프로그램이 켜져 있습니다. 근태 색인 목록을 포함해 한글 창을 모두 닫고 다시 해 주세요.")
        log("   (켜진 채로 고치면 저장이 안 되거나, 선생님이 보던 창의 내용이 덮일 수 있습니다)")
        RESULT["error"] = "hwp_running"
        return 1

    log("=" * 52)
    log(f"📋  근태 색인 목록 — {y}년 {m}월  ({'쓰기' if write else '보기만'})")
    log(f"    파일: {path.name}")
    log("=" * 52)

    # 1) 캘린더
    svc = M.get_calendar_service()
    events, bad = read_calendar(svc, M.ATTENDANCE_CALENDARS, y, m)
    missing_cal = [c for c in INCLUDE if c not in M.ATTENDANCE_CALENDARS]
    if missing_cal:
        log(f"  (캘린더 없음 — 건너뜀: {', '.join(missing_cal)})")
    for b in bad:
        RESULT["warnings"].append(f"제목 모양이 달라 읽지 못함: {b}")

    # 2) 나이스 엑셀 (기간 일수·대조)
    nd = neis_days(y, m)
    if nd is None:
        log(f"  ⚠️ 출력/{y}-{m:02d}_월별출결현황.xlsx 가 없습니다 — 교외체험 일수를 평일로 «추정»합니다")
        log(f"     (먼저 month_close.py {y}-{m:02d} --download 를 돌리면 나이스에 찍힌 날로 셉니다)")
    hol = public_holidays(svc, date(y, m, 1) - timedelta(days=40), date(y, m, 28) + timedelta(days=40))

    def count_days(e):
        days, guessed = set(), False
        cur = e["start"]
        while cur <= e["end"]:
            if cur.weekday() < 5 and cur not in hol:
                if nd is not None and (cur.year, cur.month) == (y, m):
                    if cur in nd.get((e["num"], "출석인정결석"), set()):
                        days.add(cur)
                else:
                    days.add(cur)
                    guessed = True
            cur += timedelta(days=1)
        return len(days), guessed

    # 3) 한글 파일 읽기
    kill_orphan_hwp()
    hwp = open_hwp(path)
    try:
        slots = read_rows(hwp)
        if not slots:
            log("❌ 근태 색인 목록 표를 못 찾았습니다 (머리글: 관련번호·근태상황·번호·성명·질병결석·출석인정·교외체험학습·비고)")
            RESULT["error"] = "no_table"
            return 1
        existing = []
        for s in slots:
            c = s["cells"]
            if not (c.get("date") or c.get("num") or c.get("name")):
                continue
            a, b = parse_cell_date(c.get("date"), school_year)
            existing.append({"slot": s, "start": a, "end": b, "num": c.get("num", "").strip(),
                             "name": c.get("name", "").strip(), "kind": kind_of(c), "cells": c})
        log(f"  지금 채워진 줄: {len(existing)} / {len(slots)}")

        # 4) 새로 넣을 것 고르기
        new = []
        for e in sorted(events, key=lambda e: (e["start"], int(e["num"]))):
            if e["cat"] == "질병결석":
                e["kind"] = "질병결석"
            elif e["cat"] == "출석인정결석":
                e["kind"] = "체험" if (not e["memo"] or "체험" in e["memo"]) else "출석인정"
            else:
                e["kind"] = "인정조퇴지각"
                e["end"] = e["start"]                          # 조퇴·지각은 하루짜리
            e["date_str"] = md(e["start"]) if e["start"] == e["end"] else f"{md(e['start'])}~{md(e['end'])}"
            label = f"{e['date_str']} {e['num']}번 {mask(e['name'])} {e['cat']}"
            same = [x for x in existing if x["start"] == e["start"] and x["num"] == e["num"]]
            if same:
                RESULT["skipped"].append(f"{label} — 이미 있음 (관련 {same[0]['slot']['n']})")
                continue
            overlap = [x for x in existing if x["num"] == e["num"] and x["start"] and x["end"]
                       and x["start"] <= e["end"] and e["start"] <= x["end"] and e["kind"] != "인정조퇴지각"
                       and x["kind"] in ("체험", "질병결석", "출석인정")]
            if overlap:
                RESULT["warnings"].append(f"{label} — 이미 있는 관련 {overlap[0]['slot']['n']}번과 날짜가 겹쳐 넣지 않음 (확인 필요)")
                continue
            if e["kind"] == "체험":
                e["days"], e["guessed"] = count_days(e)
                if e["days"] == 0:
                    RESULT["warnings"].append(f"{label} — 나이스 엑셀에 그 날짜의 출석인정결석이 없어 넣지 않음 (나이스 입력을 먼저 확인)")
                    continue
            # 나이스와 대조 (분류가 나이스에 찍혔는지)
            if nd is not None:
                want = {"질병결석": "질병결석", "체험": "출석인정결석", "출석인정": "출석인정결석"}.get(e["kind"], e["cat"])
                if e["start"] not in nd.get((e["num"], want), set()):
                    RESULT["warnings"].append(f"{label} — 나이스 {md(e['start'])} 에 «{want}» 로 안 보입니다 (나이스 입력 확인)")
            new.append(e)

        # 5) 날짜순으로 합치고 누적 계산
        merged = [{"src": "old", **x} for x in existing] + [{"src": "new", **x} for x in new]
        order = sorted(range(len(merged)), key=lambda i: (
            school_key(merged[i]["start"]) if merged[i]["start"] else (9999, 99, 99),
            0 if merged[i]["src"] == "old" else 1, i))
        merged = [merged[i] for i in order]
        cum = {}
        for r in merged:
            if r["kind"] != "체험":
                continue
            key = r["num"]
            if r["src"] == "old":
                try:
                    cum[key] = cum.get(key, 0) + int(re.sub(r"\D", "", r["cells"].get("days", "")) or 0)
                except ValueError:
                    pass
            else:
                cum[key] = cum.get(key, 0) + r["days"]
                r["cum"] = cum[key]
        if len(merged) > len(slots):
            log(f"❌ 줄이 모자랍니다 — {len(merged)}줄이 필요한데 표에는 {len(slots)}줄뿐입니다. 쪽(표)을 더 붙여 주세요.")
            RESULT["error"] = "full"
            return 1

        # 6) 바뀌는 줄: 처음으로 달라지는 자리부터 끝까지 다시 쓴다
        first_change = next((i for i, r in enumerate(merged) if r["src"] == "new"
                             or r["slot"]["n"] != i + 1), len(merged))
        plan = []
        for i in range(first_change, len(merged)):
            r = merged[i]
            slot = slots[i]
            cells = entry_cells(r) if r["src"] == "new" else dict(r["cells"])
            cells.pop("no", None)
            plan.append((slot, cells, r))

        log(f"\n  새로 넣을 것: {len(new)}건" + (f" · 뒤로 밀려 다시 쓰는 줄 {len(plan) - len(new)}개" if len(plan) > len(new) else ""))
        log("  관련 | 날짜 | 번호 | 이름 | 표시")
        for slot, cells, r in plan:
            if r["src"] != "new":
                continue
            show = {"질병결석": "질병결석 O", "출석인정": f"출석인정 {cells.get('excused','')}",
                    "인정조퇴지각": f"출석인정 O · 비고 {cells.get('note','')}",
                    "체험": f"교외체험 기간 {cells.get('days')}{' (추정)' if r.get('guessed') else ''} · 누적 {cells.get('cum')}"}[r["kind"]]
            log(f"  {slot['n']:>3} | {cells['date']} | {cells['num']} | {mask(cells['name'])} | {show}")
        for s in RESULT["skipped"]:
            log("  ⏭ ", s)
        for w in RESULT["warnings"]:
            log("  ⚠️ ", w)

        month_rows = [i + 1 for i, r in enumerate(merged) if r["start"] and (r["start"].year, r["start"].month) == (y, m)]
        RESULT["pages"] = sorted({(n - 1) // ROWS_PER_TABLE + 1 for n in month_rows})
        RESULT["new"] = len(new)
        if "--pages" in args or not write:
            if month_rows:
                log(f"\n  {m}월 줄: 관련 {month_rows[0]}~{month_rows[-1]} → 인쇄할 쪽 {RESULT['pages']}")
            if not write:
                log("\n  (보기만 했습니다. 맞으면 --write 를 붙여 다시 돌리세요)")
            RESULT["ok"] = True
            return 0
        if not plan:
            log("\n  ✅ 새로 넣을 것이 없습니다 — 파일은 그대로 둡니다")
            RESULT["ok"] = True
            return 0

        # 7) 백업 → 쓰기 → 저장
        bdir = OUT_DIR / "백업"
        bdir.mkdir(parents=True, exist_ok=True)
        backup = bdir / f"{path.stem}_{datetime.now():%m%d_%H%M%S}{path.suffix}"
        shutil.copy2(path, backup)
        log(f"\n  💾 고치기 전 사본: 출력/백업/{backup.name}")
        fails = []
        pages_before = hwp.PageCount
        for slot, cells, r in plan:
            b = write_row(hwp, slot, cells)
            if b:
                fails.append(f"관련 {slot['n']}: {b}")
        pages_after = hwp.PageCount
        if pages_after > pages_before:
            fails.append(f"쪽 수가 {pages_before} → {pages_after} 로 늘었습니다 — 어느 줄이 높아져 표가 다음 쪽으로 밀렸습니다. 한글로 열어 확인해 주세요")
        hwp.save_as(str(path), "HWP")
    finally:
        close_hwp(hwp)

    # 8) 저장한 파일을 «새 한글»로 다시 열어 머리글과 짝지어 확인한다
    #    (한 프로세스에서 한글을 두 번 띄우면 두 번째가 끊긴다 — 그래서 따로 띄운다)
    again = dump_rows(path)
    errors = list(fails)
    if again is None or len(again) < len(slots):
        errors.append("저장한 파일을 다시 열어 읽지 못했습니다 — 한글로 직접 열어 확인해 주세요")
        again = []
    for slot, cells, r in plan:
        if slot["n"] > len(again):
            break
        got = again[slot["n"] - 1]["cells"]
        for k, v in cells.items():
            if k in got and norm(got[k]) != norm(v):
                errors.append(f"관련 {slot['n']} {k}: 넣으려던 «{v}» ≠ 들어간 «{got[k]}»")
        marks = [k for k in ("sick", "etc", "unexcused", "excused") if got.get(k)] + (["days"] if got.get("days") else [])
        if len(marks) != 1:
            errors.append(f"관련 {slot['n']}: 표시 칸이 {len(marks)}개 ({marks}) — 하나여야 합니다")
    if errors:
        log("\n  ❌ 다시 읽어 보니 어긋난 칸이 있습니다:")
        for x in errors[:20]:
            log("    ", x)
        log(f"     고치기 전 사본은 출력/백업/{backup.name} 에 있습니다")
        RESULT["error"] = "verify"
        RESULT["verify_errors"] = errors[:20]
    else:
        log(f"\n  ✅ 저장하고 다시 읽어 확인했습니다 — {len(plan)}줄 모두 머리글과 맞습니다")
        RESULT["ok"] = True
    return 0 if RESULT["ok"] else 1


if __name__ == "__main__":
    if len(sys.argv) == 3 and sys.argv[1] == "--_dump":
        _dump(sys.argv[2])
        sys.exit(0)
    code = 1
    try:
        code = main()
    except Exception as e:
        import traceback
        traceback.print_exc()
        RESULT["error"] = str(e)[:200]
    print("RESULT_JSON " + json.dumps(RESULT, ensure_ascii=False))
    sys.exit(code)
